import re
import datetime as _dt
from io import BytesIO

import pandas as pd
import streamlit as st

st.set_page_config(page_title="Header Fixer", layout="wide")

# ---------------------------------------------------------------------------
# Standard/target headers for each report type. Edit these lists directly if
# a standard header ever changes.
# ---------------------------------------------------------------------------
DAILY_REMARK_HEADER = [
    "S.No", "Date", "Time", "Debtor", "Account No.", "Card No.", "Service No.",
    "DPD", "Call Status", "Status", "Remark", "Remark By", "Remark Type",
    "Field Visit Date", "Collector", "Client", "Product Description",
    "Product Type", "Batch No", "Account Type", "Relation", "PTP Amount",
    "PTP Date", "Next Call", "Claim Paid Amount", "Claim Paid Date",
    "Dialed Number", "Days Past Write Off", "Balance", "Contact Type",
    "Cycle", "Old IC", "I.C Issue Date", "Bank Code", "Over Limit Amount",
    "Min Payment", "Due Date", "Monthly Installment", "30 Days", "MIA",
    "Area", "Call Duration", "Talk Time Duration", "Debtor ID",
    "Black Case No.", "Red Case No.", "Court Name", "Lawyer", "Legal Stage",
    "Legal Status", "Next Legal Follow up",
]

AUTOSTAT_HEADER = [
    "CH CODE", "ACCOUNT NUMBER", "STATUS CODE", "REMARKS", "REMARKS BY",
    "REMARKS DATE", "PTP DATE", "PTP AMOUNT",
]

# Autostat date/time columns need special reformatting (the system expects
# these exact patterns to recognize the file): REMARKS DATE -> MM/DD/YYYY
# HH:MM:SS (24-hour), PTP DATE -> MM/DD/YYYY (date only).
AUTOSTAT_DATETIME_COLS = {"REMARKS DATE": False, "PTP DATE": True}  # value = date_only?

# ---------------------------------------------------------------------------
# DRR cleaning rules. Edit these lists if the rules change.
# Status is matched EXACTLY (case-insensitive), so "Field" will not hit
# something like "Field Visit". Remark is matched by "contains".
# ---------------------------------------------------------------------------
BASE_BAD_STATUSES = ["BP", "New", "Reactive", "Aborted", "Locked", "Unlocked", "SMS Failed"]
BAD_REMARK_PHRASES = ["New Assignment", "Updates when case", "System Auto PD", "New Contact Details"]

CLEANING_PROFILES = {
    "ECMS cleaning": BASE_BAD_STATUSES + ["Field"],
    "Status DRR cleaning (strict filters)": BASE_BAD_STATUSES,
}

# If the source system ever renames a column, map normalized source name ->
# normalized target name here, e.g. {"acctno": "accountno"}.
HEADER_ALIASES = {}

DRR_DATE_FMT_XLSX = "DD-MM-YYYY"
DRR_TIME_FMT_XLSX = "hh:mm:ss am/pm"

REPORT_CONFIGS = {
    "Daily Remark": {"target_header": DAILY_REMARK_HEADER, "datetime_cols": {}, "is_drr": True},
    "Autostat": {"target_header": AUTOSTAT_HEADER, "datetime_cols": AUTOSTAT_DATETIME_COLS, "is_drr": False},
}


# ---------------------------------------------------------------------------
# Reading files
# ---------------------------------------------------------------------------
@st.cache_data(show_spinner="Reading file...")
def load_file(file_bytes: bytes, name: str) -> pd.DataFrame:
    """Read a CSV or Excel file with native cell values preserved.

    Excel is read with dtype=object so long account/card numbers keep their
    digits and real date/time cells stay real dates/times (we format them
    ourselves later).
    """
    buf = BytesIO(file_bytes)
    if name.lower().endswith(".csv"):
        return pd.read_csv(buf, dtype=str, keep_default_na=False)
    return pd.read_excel(buf, dtype=object, keep_default_na=False)


def _format_cell(val):
    """Turn any cell value into clean text."""
    if val is None:
        return ""
    try:
        if pd.isna(val):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(val, pd.Timestamp):
        val = val.to_pydatetime()
    if isinstance(val, _dt.datetime):
        if val.time() == _dt.time(0, 0):
            return val.strftime("%d/%m/%Y")
        return val.strftime("%d/%m/%Y %I:%M:%S %p")
    if isinstance(val, _dt.date):
        return val.strftime("%d/%m/%Y")
    if isinstance(val, _dt.time):
        return val.strftime("%I:%M:%S %p")
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return str(val)
    return str(val).strip()


def normalize_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    return df.apply(lambda col: col.map(_format_cell))


# ---------------------------------------------------------------------------
# Autostat date/time handling (unchanged behaviour)
# ---------------------------------------------------------------------------
_TRAILING_AMPM_RE = re.compile(r"\s*(AM|PM)\s*$", re.IGNORECASE)


def parse_autostat_datetime(raw, date_only: bool = False) -> str:
    if raw is None:
        return ""
    s = str(raw).strip()
    if not s or s.lower() == "nan":
        return ""
    s = re.sub(r"\s+", " ", s)
    s = _TRAILING_AMPM_RE.sub("", s)

    formats = ["%m/%d/%Y"] if date_only else [
        "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y",
    ]
    for fmt in formats:
        try:
            dt = _dt.datetime.strptime(s, fmt)
            return dt.strftime("%m/%d/%Y") if date_only else dt.strftime("%m/%d/%Y %H:%M:%S")
        except ValueError:
            continue
    try:
        dt = pd.to_datetime(s)
        return dt.strftime("%m/%d/%Y") if date_only else dt.strftime("%m/%d/%Y %H:%M:%S")
    except Exception:
        return s


def apply_datetime_formatting(df: pd.DataFrame, datetime_cols: dict) -> pd.DataFrame:
    for col, date_only in datetime_cols.items():
        if col in df.columns:
            df[col] = df[col].map(lambda v: parse_autostat_datetime(v, date_only=date_only))
    return df


# ---------------------------------------------------------------------------
# Header alignment
# ---------------------------------------------------------------------------
def _norm(s) -> str:
    """Case, space and punctuation-insensitive key: 'Account No.' == 'account no'."""
    key = re.sub(r"[^a-z0-9]", "", str(s).lower())
    return HEADER_ALIASES.get(key, key)


def fix_header(df: pd.DataFrame, target_header: list):
    """Reorder/select df's columns to match target_header.
    Returns (fixed_df, missing_cols, extra_cols).
    """
    col_pos = {}
    for i, c in enumerate(df.columns):
        col_pos.setdefault(_norm(c), i)

    data, missing_cols = {}, []
    for target_col in target_header:
        pos = col_pos.get(_norm(target_col))
        if pos is not None:
            data[target_col] = df.iloc[:, pos].values
        else:
            data[target_col] = [""] * len(df)
            missing_cols.append(target_col)

    used_keys = {_norm(t) for t in target_header}
    extra_cols = [c for c in df.columns if _norm(c) not in used_keys]
    return pd.DataFrame(data), missing_cols, extra_cols


# ---------------------------------------------------------------------------
# DRR: date / time parsing, cleaning, sorting
# ---------------------------------------------------------------------------
def parse_date_value(v):
    """Any Date cell -> pandas Timestamp (date only) or NaT. Text dates are
    read day-first (dd/mm/yyyy)."""
    if v is None or v is pd.NaT:
        return pd.NaT
    if isinstance(v, _dt.datetime) or isinstance(v, _dt.date):
        return pd.Timestamp(v).normalize()
    s = str(v).strip()
    if not s or s.lower() in ("nan", "nat"):
        return pd.NaT
    s = s.split(" ")[0]  # drop any time part
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%Y/%m/%d", "%d/%m/%y"):
        try:
            return pd.Timestamp(_dt.datetime.strptime(s, fmt))
        except ValueError:
            continue
    return pd.to_datetime(s, dayfirst=True, errors="coerce")


def parse_time_value(v):
    """Any Time cell -> datetime.time or None."""
    if v is None or v is pd.NaT:
        return None
    if isinstance(v, _dt.datetime):
        return v.time().replace(microsecond=0)
    if isinstance(v, _dt.time):
        return v.replace(microsecond=0)
    s = re.sub(r"\s+", " ", str(v).strip())
    if not s or s.lower() in ("nan", "nat"):
        return None
    for fmt in ("%I:%M:%S %p", "%I:%M %p", "%H:%M:%S", "%H:%M"):
        try:
            return _dt.datetime.strptime(s, fmt).time()
        except ValueError:
            continue
    stripped = _TRAILING_AMPM_RE.sub("", s)  # e.g. '18:00:00 PM'
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return _dt.datetime.strptime(stripped, fmt).time()
        except ValueError:
            continue
    return None


def find_rows_to_clean(df: pd.DataFrame, bad_statuses: list):
    """Return (mask, status_hits, remark_hits) for rows matching the rules."""
    status_norm = df["Status"].astype(str).str.strip().str.lower()
    remark_norm = df["Remark"].astype(str).str.lower()

    mask = pd.Series(False, index=df.index)
    status_hits, remark_hits = {}, {}
    for s in bad_statuses:
        hit = status_norm == s.lower()
        if hit.any():
            status_hits[s] = int(hit.sum())
            mask |= hit
    for phrase in BAD_REMARK_PHRASES:
        hit = remark_norm.str.contains(phrase.lower(), regex=False)
        if hit.any():
            remark_hits[phrase] = int(hit.sum())
            mask |= hit
    return mask, status_hits, remark_hits


def to_excel_bytes(df: pd.DataFrame, col_formats: dict = None) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="sheet")
    if not col_formats:
        return buffer.getvalue()

    # Apply Excel number formats after the sheet is written.
    from openpyxl import load_workbook
    buffer.seek(0)
    wb = load_workbook(buffer)
    ws = wb.active
    ws.title = "sheet"
    for col_name, fmt in col_formats.items():
        if col_name in df.columns:
            idx = list(df.columns).index(col_name) + 1
            for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                row[0].number_format = fmt
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def process_drr(merged: pd.DataFrame, drop_dupes: bool = True):
    """Parse Date/Time, text-format everything else, optionally drop exact
    duplicate rows. Returns (df, dup_count)."""
    df = merged.copy()
    dates = pd.to_datetime(df["Date"].map(parse_date_value))
    times = df["Time"].map(parse_time_value)

    text_cols = [c for c in df.columns if c not in ("Date", "Time")]
    df[text_cols] = df[text_cols].apply(lambda col: col.map(_format_cell))
    df["Date"] = dates
    df["Time"] = times

    dupes = df.duplicated(keep="first")
    if drop_dupes:
        df = df[~dupes].reset_index(drop=True)
    return df, int(dupes.sum())


def sort_newest_first(df: pd.DataFrame) -> pd.DataFrame:
    secs = df["Time"].map(lambda t: 0 if t is None else t.hour * 3600 + t.minute * 60 + t.second)
    stamp = df["Date"] + pd.to_timedelta(secs, unit="s")
    order = stamp.sort_values(ascending=False, kind="mergesort", na_position="last").index
    return df.loc[order].reset_index(drop=True)


def drr_preview_text(df: pd.DataFrame) -> pd.DataFrame:
    """Date as DD-MM-YYYY text, Time as hh:mm:ss am/pm text."""
    out = df.copy()
    out["Date"] = out["Date"].dt.strftime("%d-%m-%Y").fillna("")
    out["Time"] = out["Time"].map(lambda t: "" if t is None else t.strftime("%I:%M:%S %p").lower())
    return out


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------
st.title("📋 Header Fixer")
st.caption(
    "Drag in one or more downloaded files. They get merged, aligned to your standard "
    "header, and (for Daily Remark) cleaned and sorted newest to oldest."
)

report_type = st.radio("Which file are you fixing?", list(REPORT_CONFIGS.keys()), horizontal=True)
config = REPORT_CONFIGS[report_type]
target_header = config["target_header"]
datetime_cols = config["datetime_cols"]
is_drr = config["is_drr"]

with st.expander(f"Standard header currently in use for {report_type} (edit app.py to change it)"):
    st.write(f"{len(target_header)} columns:")
    st.code(", ".join(target_header))

uploaded_files = st.file_uploader(
    "Upload the data file(s) (.csv or .xlsx). You can drag in several at once.",
    type=["csv", "xlsx", "xls"],
    accept_multiple_files=True,
    key=report_type,
)

if not uploaded_files:
    st.info("Upload one or more files above to get started.")
    st.stop()

# ----- load + align each file, then merge ---------------------------------
frames, summary, missing_by_file, extra_cols_all = [], [], {}, []
for f in uploaded_files:
    try:
        raw = load_file(f.getvalue(), f.name)
    except Exception as e:
        st.error(f"Couldn't read {f.name}: {e}")
        st.stop()

    summary.append({"File": f.name, "Rows": raw.shape[0], "Columns": raw.shape[1]})

    if not is_drr:
        raw = normalize_dataframe(raw)
    fixed, missing, extra = fix_header(raw, target_header)
    if not is_drr and datetime_cols:
        fixed = apply_datetime_formatting(fixed, datetime_cols)

    frames.append(fixed)
    if missing:
        missing_by_file[f.name] = missing
    for c in extra:
        if c not in extra_cols_all:
            extra_cols_all.append(c)

merged = pd.concat(frames, ignore_index=True)
st.success(f"Loaded {len(uploaded_files)} file(s), {len(merged)} rows in total")
if len(uploaded_files) > 1:
    st.dataframe(pd.DataFrame(summary), use_container_width=True, hide_index=True)

col1, col2 = st.columns(2)
with col1:
    if missing_by_file:
        lines = [f"- **{name}**: {', '.join(cols)}" for name, cols in missing_by_file.items()]
        st.warning("⚠️ Standard columns NOT found (left blank):\n\n" + "\n".join(lines))
    else:
        st.info("✅ All standard columns were found in every file.")
with col2:
    if extra_cols_all:
        st.info(
            f"ℹ️ {len(extra_cols_all)} extra column(s) dropped (not in your standard header):\n\n- "
            + "\n- ".join(map(str, extra_cols_all))
        )
    else:
        st.info("No extra columns found.")

# ----- Autostat: simple path ----------------------------------------------
if not is_drr:
    out_df = merged
    st.subheader("Preview (first 20 rows)")
    st.dataframe(out_df.head(20), use_container_width=True)
    st.download_button(
        label="⬇️ Download fixed file (.xlsx)",
        data=to_excel_bytes(out_df),
        file_name=f"{report_type.replace(' ', '_')}_Fixed.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    st.stop()

# ----- Daily Remark path ---------------------------------------------------
st.divider()
st.subheader("Daily Remark options")
opt1, opt2 = st.columns(2)
with opt1:
    profile = st.radio("Cleaning profile", list(CLEANING_PROFILES.keys()))
with opt2:
    drop_dupes = st.checkbox("Remove exact duplicate rows (useful when pulled day by day and files overlap)", value=True)
    as_text = st.checkbox("Write Date/Time as plain text instead of real Excel dates", value=False)

df, dup_count = process_drr(merged, drop_dupes=drop_dupes)
if dup_count:
    if drop_dupes:
        st.info(f"Removed {dup_count} exact duplicate row(s) across the uploaded files.")
    else:
        st.warning(f"{dup_count} exact duplicate row(s) found and kept.")

# ----- cleaning: ask first, then delete -------------------------------------
mask, status_hits, remark_hits = find_rows_to_clean(df, CLEANING_PROFILES[profile])
if mask.any():
    st.warning(f"Found {int(mask.sum())} row(s) that the {profile} rules delete.")
    c1, c2 = st.columns(2)
    with c1:
        if status_hits:
            st.write("**Status column**")
            st.dataframe(
                pd.DataFrame({"Status": list(status_hits), "Rows": list(status_hits.values())}),
                hide_index=True, use_container_width=True,
            )
    with c2:
        if remark_hits:
            st.write("**Remark column (contains)**")
            st.dataframe(
                pd.DataFrame({"Remark contains": list(remark_hits), "Rows": list(remark_hits.values())}),
                hide_index=True, use_container_width=True,
            )
    with st.expander("See the rows that would be deleted"):
        st.dataframe(drr_preview_text(df[mask]), use_container_width=True)

    sig = "|".join(f"{f.name}:{f.size}" for f in uploaded_files)
    decision = st.radio(
        "Delete these rows?",
        ["Yes, delete them", "No, keep them"],
        index=None,
        horizontal=True,
        key=f"clean_{profile}_{sig}",
    )
    if decision is None:
        st.stop()  # wait for the answer before showing / offering the file
    if decision.startswith("Yes"):
        df = df[~mask].reset_index(drop=True)
        st.success("Rows deleted.")
else:
    st.info("✅ No rows matched the cleaning rules.")

# ----- sort newest -> oldest -------------------------------------------------
df = sort_newest_first(df)

bad_dates = int(df["Date"].isna().sum())
if bad_dates:
    st.warning(f"{bad_dates} row(s) have no readable Date and were placed at the bottom.")

valid_dates = df["Date"].dropna()
if not valid_dates.empty:
    lo, hi = valid_dates.min(), valid_dates.max()
    st.caption(f"Covers {lo:%d-%m-%Y} to {hi:%d-%m-%Y} · {len(df)} rows · newest first")
    present = set(valid_dates.dt.normalize())
    gaps = [d for d in pd.date_range(lo, hi) if d not in present]
    if gaps:
        st.warning("No rows found for these date(s) in the range: " + ", ".join(f"{d:%d-%m-%Y}" for d in gaps))

# ----- preview + download ----------------------------------------------------
preview = drr_preview_text(df)
st.subheader("Preview (first 20 rows)")
st.dataframe(preview.head(20), use_container_width=True)

if as_text:
    excel_bytes = to_excel_bytes(preview)
else:
    excel_bytes = to_excel_bytes(df, {"Date": DRR_DATE_FMT_XLSX, "Time": DRR_TIME_FMT_XLSX})

st.download_button(
    label="⬇️ Download fixed file (.xlsx)",
    data=excel_bytes,
    file_name=f"{report_type.replace(' ', '_')}_Fixed.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)
